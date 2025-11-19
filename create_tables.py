import os
import pandas as pd
from docx import Document
from docx.shared import Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_col_width(cell, width_in_inches):
    """Ustawia szerokość kolumny i włącza zawijanie tekstu"""
    cell.width = Inches(width_in_inches)
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcW = OxmlElement('w:tcW')
    tcW.set(qn('w:w'), str(int(width_in_inches * 1440)))  # twips
    tcW.set(qn('w:type'), 'dxa')
    tcPr.append(tcW)
    noWrap = tcPr.find(qn('w:noWrap'))
    if noWrap is not None:
        tcPr.remove(noWrap)

word_filename = "GA_results.docx"
if os.path.exists(word_filename):
    doc = Document(word_filename)
else:
    doc = Document()
    doc.add_heading("Wyniki Algorytmu Genetycznego", level=1)

headers_main = ['Pc','Pm',"N","T","Najlepsza wartość",'Najniższa wartość','Czas wykonywania']
widths_main = [0.7, 0.7, 0.7, 0.7, 1.0, 1.0, 1.0]

csv_files_final = sorted([f for f in os.listdir(".") if f.startswith("final_results") and f.endswith(".csv")])

combination_number = 1
for idx, csv_file in enumerate(csv_files_final):
    if idx % 9 == 0:
        doc.add_paragraph(f"Kombinacja {combination_number}", style='Heading 2')
        combination_number += 1

    doc.add_paragraph(f"Wyniki z pliku {csv_file}", style='Heading 3')
    df = pd.read_csv(csv_file)

    table_main = doc.add_table(rows=1, cols=len(headers_main))
    table_main.style = 'Table Grid'
    for i, header in enumerate(headers_main):
        table_main.rows[0].cells[i].text = header
        set_col_width(table_main.rows[0].cells[i], widths_main[i])

    for _, row in df.iterrows():
        new_cells = table_main.add_row().cells
        new_cells[0].text = str(row['Pc'])
        new_cells[1].text = str(row['Pm'])
        new_cells[2].text = str(row['N'])
        new_cells[3].text = str(row['T'])
        new_cells[4].text = str(row['Best_value'])
        new_cells[5].text = str(row['Worst_value'])
        new_cells[6].text = str(row['Time_exec'])
        for cell, width in zip(new_cells, widths_main):
            set_col_width(cell, width)

    doc.add_paragraph("Najlepsze itemy:", style='Normal')
    table_items = doc.add_table(rows=1, cols=1)
    table_items.style = 'Table Grid'
    table_items.rows[0].cells[0].text = "Best_items"
    set_col_width(table_items.rows[0].cells[0], 5.0)

    for _, row in df.iterrows():
        new_cells = table_items.add_row().cells
        new_cells[0].text = str(row['Best_items'])
        set_col_width(new_cells[0], 5.0)

doc.save(word_filename)
print(f"Wyniki zapisane w {word_filename}")
